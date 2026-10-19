import json,sqlite3
from pathlib import Path
import pytest
from openpyxl import load_workbook
from job_agent.database import now
from job_agent.jobqueue import COMMON_SHEETS,JobQueueWorkbook
from job_agent.m2d import GmailExportAdapter,M2DEngine,canonical_url,classify_readiness,headcount_state,stable_job_key
from job_agent.m2d_models import CoverageStatus,EmailExportMessage,OperationalFacts,Readiness,TriState
from job_agent.m2d_sources import PRIORITY_EMPLOYERS

def engine(app):return app.state.m2d

def test_source_matrix_ownership_and_honest_coverage(app):
    m=engine(app);sources=m.sources();names={x['display_name'] for x in sources}
    assert {'LinkedIn Jobs','Naukri','Indeed','Glassdoor','Foundit','Remote OK','We Work Remotely','Received Gmail job messages'}<=names
    assert {x['stream'] for x in sources}=={'MAJOR','SPECIALIST','REMOTE','GMAIL','PIPELINE'}
    gmail=next(x for x in sources if x['source_id']=='gmail');assert gmail['implementation_state']=='ACCESS_BLOCKED' and gmail['coverage_status']=='ACCESS_BLOCKED'
    coverage=m.coverage();assert coverage['configured']==len(sources)-1 and coverage['contracts']==1 and coverage['attempted']>=1 and coverage['claim']=='configured_discovery_sources_only'
    assert all(x['reference_time_ist'] for x in sources)

def test_readiness_unknowns_thresholds_and_recipient_independence():
    facts=OperationalFacts(active='PASS',india_eligible='PASS',headcount='UNKNOWN',direct_permanent_payroll='UNKNOWN',mandatory_requirements='PASS')
    state,missing,_=classify_readiness(facts);assert state==Readiness.READY_FOR_PORTAL_REVIEW and 'headcount' in missing and 'verified_recipient' not in missing
    complete=facts.model_copy(update={'headcount':TriState.PASS,'direct_permanent_payroll':TriState.PASS})
    assert classify_readiness(complete)[0]==Readiness.FULLY_VERIFIED_APPLICATION_READY
    assert classify_readiness(complete.model_copy(update={'verified_recipient':TriState.PASS}))[0]==Readiness.EMAIL_DRAFT_READY
    assert classify_readiness(complete.model_copy(update={'india_eligible':TriState.FAIL}))[0]==Readiness.EXCLUDED
    assert headcount_state(None,'remote')==TriState.UNKNOWN and headcount_state(500,'remote')==TriState.PASS and headcount_state(999,'hybrid')==TriState.FAIL and headcount_state(1000,'hybrid')==TriState.PASS

def test_stable_identity_uses_specific_vacancy_and_preserves_distinct_requisitions():
    first=stable_job_key(url='https://EXAMPLE.com/jobs/1?utm_source=x',source_id='a',provider_id='1',company='Acme',title='Engineer')
    assert first==stable_job_key(url='https://example.com/jobs/1',source_id='b',provider_id='other',company='Different',title='Other')
    assert first!=stable_job_key(url='https://example.com/jobs/2',source_id='a',provider_id='2',company='Acme',title='Engineer')
    assert canonical_url('javascript:alert(1)') is None

def test_queue_fairness_recovery_bounded_continuation_and_attempts(app):
    m=engine(app)
    for stream,source in [('MAJOR','linkedin'),('REMOTE','remoteok'),('GMAIL','gmail'),('SPECIALIST','hirist')]:m.enqueue(kind='DISCOVERY',source_id=source,stream=stream,category='SWEEP',task_id='task-'+stream)
    claimed=m.claim(limit=4,lease_seconds=1);assert {x['stream'] for x in claimed}=={'MAJOR','REMOTE','GMAIL','SPECIALIST'}
    assert m.finish_task('task-MAJOR',status='PARTIAL',checkpoint={'page':1},continuation={'page':2})
    with m.db.transaction() as conn:conn.execute("UPDATE m2d_tasks SET lease_until='2000-01-01' WHERE task_id='task-REMOTE'")
    assert m.recover_expired()>=1
    continuation=next(x for x in m.tasks() if x['task_id'] not in {'task-MAJOR','task-REMOTE','task-GMAIL','task-SPECIALIST'})
    assert continuation['status']=='PENDING' and 'page' in continuation['payload']

def test_week_plan_rotates_all_employers_and_is_idempotent(app):
    m=engine(app);first=m.plan_week('2026-10-14T00:00:00Z');second=m.plan_week('2026-10-14T00:00:00Z')
    assert first==second
    tasks=m.tasks();groups=[json.loads(x['payload'])['employers'] for x in tasks if x['category']=='EMPLOYER_GROUP']
    assert len({name for group in groups for name in group})==len(PRIORITY_EMPLOYERS)
    assert all(x['next_due'] for x in m.sources())

def test_coverage_requires_bounded_exhaustion_and_reports_separate_counts(app):
    m=engine(app)
    with pytest.raises(ValueError,match='exhaustive_requires'):m.record_coverage('remoteok',CoverageStatus.EXHAUSTIVELY_SEARCHED)
    assert m.record_coverage('remoteok',CoverageStatus.EXHAUSTIVELY_SEARCHED,scope={'query':'frontend','pages':2},cursor={'complete':True},exhausted=True)
    result=m.coverage();assert result['exhaustive']==1 and result['meaningful']>=1 and result['attempted']>=result['meaningful']

def message(mid,body,**changes):
    raw=dict(message_id=mid,thread_id='thread',received_at='2026-10-14T04:00:00Z',labels=['INBOX'],mailbox='received',subject='Jobs',body=body,cursor=mid);raw.update(changes);return EmailExportMessage(**raw)

def test_gmail_export_received_only_multi_job_pagination_and_idempotence(app):
    m=engine(app);adapter=GmailExportAdapter(m);body='Job: React Lead\nCompany: Acme\nURL: https://example.com/jobs/1\n\nRole: AI Engineer\nCompany: Beta\nLink: https://example.com/jobs/2'
    messages=[message('m1',body),message('m2','Course: improve your profile',subject='Career course'),message('m3',body,mailbox='sent')]
    first=adapter.ingest(messages,limit=2,interval_end='2026-10-14T05:00:00Z');assert first['vacancies']==2 and first['status']=='PARTIAL' and first['remaining']==1
    assert adapter.ingest([messages[0]])['vacancies']==0
    newer=adapter.ingest([message('m4','Role: Platform Lead\nCompany: Gamma\nURL: https://example.com/jobs/3')]);assert newer['vacancies']==1
    with m.db.transaction() as conn:
        assert conn.execute('select count(1) from gmail_intake').fetchone()[0]==3
        assert conn.execute("select count(distinct job_key) from job_observations where stream='GMAIL'").fetchone()[0]==3
        assert 'React Lead' not in str(conn.execute('select * from gmail_intake').fetchall())

def test_snapshot_immutable_consumer_idempotent_and_stale_retained(app):
    m=engine(app);sid=m.snapshot(run_id='run',stream='MAJOR',status='PARTIAL',started_at=now(),counts={'pending':2})
    assert m.consume('pipeline',sid) and not m.consume('pipeline',sid)
    with m.db.transaction() as conn:
        with pytest.raises(sqlite3.IntegrityError,match='immutable_snapshot'):conn.execute("update m2d_snapshots set completion_status='COMPLETED' where snapshot_id=?",(sid,))
    assert m.summary()['status']=='PARTIAL'

def test_legacy_run_reconciliation_is_idempotent_and_bounded(app):
    m=engine(app)
    with m.db.transaction() as conn:
        conn.execute("INSERT INTO sources(id,provider,company,board,region,enabled,checkpoint) VALUES('fixture-gh','greenhouse','Fixture','fixture','global',1,?)",(json.dumps({'complete':True,'pages':1}),))
    run={'id':'legacy-run','status':'COMPLETED','started_at':'2026-10-14T00:00:00Z','metrics':{'new_jobs':2,'updated_jobs':0},'sources':[{'source_id':'fixture-gh','status':'SUCCESS','error':None}]}
    snapshot=m.reconcile_discovery_run(run);assert snapshot and m.reconcile_discovery_run(run) is None
    source=next(x for x in m.sources() if x['source_id']=='greenhouse');assert source['coverage_status']=='EXHAUSTIVELY_SEARCHED' and source['attempts']==1

def test_reconciliation_aggregates_provider_rejections_and_queues_verification(app):
    m=engine(app)
    with m.db.transaction() as conn:
        for suffix in ('a','b'):conn.execute("INSERT INTO sources(id,provider,company,board,region,enabled,checkpoint) VALUES(?,?,?,?,?,?,?)",(f'gh-{suffix}','greenhouse','Fixture',suffix,'global',1,json.dumps({'complete':True})))
    run={'id':'partial-run','status':'PARTIAL','started_at':'2026-10-14T00:00:00Z','metrics':{'new_jobs':1,'updated_jobs':0},'sources':[
        {'source_id':'gh-a','status':'SUCCESS','error':None,'metrics':{'rejected_records':0}},
        {'source_id':'gh-b','status':'PARTIAL','error':None,'metrics':{'rejected_records':7}}]}
    assert m.reconcile_discovery_run(run)
    source=next(x for x in m.sources() if x['source_id']=='greenhouse');assert source['backlog']==7 and source['attempts']==1 and source['blocker']=='rejected_records:7'
    key=m.observe(stream='MAJOR',source_id='fixture',provider_id='verify',url='https://example.com/verify',company='Example',title='Lead')
    m.set_operational_state(key,OperationalFacts(active='PASS'),strong_fit=True)
    with m.db.transaction() as conn:conn.execute("UPDATE job_operational_state SET job_id=NULL WHERE job_key=?",(key,))
    m.seed_verification_tasks();assert any(x['job_key']==key for x in m.tasks('PENDING'))

def test_profile_conflicts_are_reported_without_override(app):
    m=engine(app);background=app.state.discovery.registry.background().model_copy(update={'total_years':6})
    from job_agent.candidate import CandidateProfile
    profile=CandidateProfile(total_experience_years=8,preferred_locations=['noida'])
    conflicts=m.profile_conflicts(background,profile);assert {x['field'] for x in conflicts}=={'search_background.total_years','candidate.total_experience_years','candidate.preferred_locations'}
    assert background.total_years==6 and profile.total_experience_years==8

def test_jobqueue_v1_sheets_hyperlinks_roundtrip_and_status_claim_is_not_readiness(app,tmp_path):
    m=engine(app);key=m.observe(stream='MAJOR',source_id='fixture',provider_id='req-1',url='https://example.com/jobs/req-1',company='Example',title='Frontend Lead',evidence={'synthetic':True})
    m.set_operational_state(key,OperationalFacts(active='PASS',india_eligible='PASS',headcount='UNKNOWN',direct_permanent_payroll='UNKNOWN',mandatory_requirements='PASS'))
    result=JobQueueWorkbook(m).export('MAJOR',tmp_path,run_id='fixture-run');book=load_workbook(result['path'])
    assert set(COMMON_SHEETS)|{'Employer Careers Registry','Application Review Shortlist'}<=set(book.sheetnames)
    headers=[c.value for c in book['Jobs'][1]];link_col=headers.index('ExactVacancyURL')+1;assert book['Jobs'].cell(2,link_col).hyperlink.target=='https://example.com/jobs/req-1'
    imported=JobQueueWorkbook(m).import_workbook(result['path'],'MAJOR');assert imported['rows_imported']==1 and not imported['duplicate']
    assert JobQueueWorkbook(m).import_workbook(result['path'],'MAJOR')['duplicate']
    with m.db.transaction() as conn:
        assert conn.execute("select readiness from job_operational_state where job_key=?",(key,)).fetchone()[0]=='READY_FOR_PORTAL_REVIEW'
        assert conn.execute("select imported_status_claim from m2d_imports where import_id=?",(imported['import_id'],)).fetchone()[0] is None
    assert (tmp_path/'manual-transfer-manifest.json').read_text().find('NOT_UPLOADED')>0
    consolidated=JobQueueWorkbook(m).export('CONSOLIDATED',tmp_path,run_id='consolidated-fixture');assert load_workbook(consolidated['path'])['Application Review Shortlist'].max_row>=2

def test_m2d_api_promising_visible_and_no_m3_actions(client):
    summary=client.get('/api/m2d/summary');assert summary.status_code==200 and summary.json()['coverage']['configured']>20
    assert client.get('/api/m2d/sources').status_code==200 and client.get('/api/m2d/tasks').status_code==200
    assert client.post('/api/send-email',json={}).status_code==404 and client.post('/api/gmail/drafts',json={}).status_code==404

def test_v8_and_prior_backup_existing_rows_preserved(app):
    db=app.state.service.db
    with db.transaction() as conn:assert conn.execute('pragma user_version').fetchone()[0]==17
    assert db.path.with_suffix(db.path.suffix+'.pre-v7.bak').exists() and db.path.with_suffix(db.path.suffix+'.pre-v8.bak').exists()
