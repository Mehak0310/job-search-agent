import httpx
import pytest

from job_agent.database import Database
from job_agent.discovery import Discovery
from job_agent.discovery_models import DiscoveryConfig, DiscoverySettings, SourceConfig, metrics
from job_agent.m2d import GmailExportAdapter, canonical_url
from job_agent.public_connectors import PublicReader, RequestBudget, ProviderError
from job_agent.registry import Registry
from job_agent.smartrecruiters import SmartRecruitersConnector
from job_agent.amazon_connector import AmazonConnector
from job_agent.models import Profile
from job_agent.public_connectors import Batch
from job_agent.m2d import M2DEngine
from job_agent.m2d_models import EmailExportMessage
from job_agent.coverage_report import report
from job_agent.gmail_reextract import reextract_interval
from job_agent.discovery_store import persist_batch
from job_agent.service import JobService
from argparse import Namespace
from job_agent import release


def source():
    return SourceConfig(id="nagarro-smartrecruiters", provider="smartrecruiters", company="Nagarro", board="Nagarro1", enabled=True)


def posting(identifier="1"):
    return {"id":identifier,"name":"Senior React Engineer","company":{"identifier":"Nagarro1","name":"Nagarro"},
        "location":{"city":"Noida","country":"in","hybrid":True},"releasedDate":"2026-10-15T01:00:00Z"}


def test_pagination_budget_resume_and_distinct_requisitions(tmp_path):
    calls=[]
    state={'total':3}
    def handler(request):
        if 'offset' not in request.url.params:
            return httpx.Response(200,json=posting(request.url.path.rsplit('/',1)[-1])|{'jobAd':{'sections':{'qualifications':{'title':'Qualifications','text':'React required.'}}}})
        offset=int(request.url.params['offset']);calls.append(offset)
        return httpx.Response(200,json={"offset":offset,"totalFound":state['total'],"content":[posting(str(offset+1))]})
    db=Database(tmp_path/'state.sqlite3')
    cfg=DiscoveryConfig(sources=[source()],settings=DiscoverySettings(max_requests=2,page_size=1,cooldown_seconds=0,retry_limit=0))
    Registry(db).sync(cfg)
    discovery=Discovery(db,cfg.settings,lambda:httpx.Client(transport=httpx.MockTransport(handler)))
    try:
        first=discovery.execute(discovery.claim())
        assert first['status']=='PARTIAL' and first['metrics']['new_jobs']==2
        checkpoint=Registry(db).list()[0]['checkpoint']
        assert checkpoint['next_offset']==2 and checkpoint['termination_reason']=='request_budget_exhausted'
        second=discovery.execute(discovery.claim())
        assert second['status']=='COMPLETED' and second['metrics']['new_jobs']==1
        with db.transaction() as conn:
            assert conn.execute('select count(*) from jobs').fetchone()[0]==3
            assert conn.execute("select count(*) from jobs where activity_status='ACTIVE'").fetchone()[0]==3
            assert conn.execute('pragma foreign_key_check').fetchall()==[]
        assert calls==[0,1,2]
        assert Registry(db).list()[0]['checkpoint']['scan_seen_ids']==['1','2','3']
        state['total']=2
        discovery.execute(discovery.claim());discovery.execute(discovery.claim())
        with db.transaction() as conn:
            assert conn.execute("select activity_status from jobs j join source_references r on r.job_id=j.id where r.source_key='3'").fetchone()[0]=='CLOSED'
    finally:discovery.close()


@pytest.mark.parametrize('payload,code',[
    ({'offset':0,'totalFound':3,'content':[]},'invalid_pagination'),
    ({'offset':1,'totalFound':1,'content':[posting()]},'invalid_pagination'),
    ({'content':{}},'invalid_board_response'),
    ({'offset':0,'totalFound':1,'content':[{'id':'../evil'}]},'invalid_source_record'),
])
def test_invalid_pages(payload,code):
    client=httpx.Client(transport=httpx.MockTransport(lambda _:httpx.Response(200,json=payload)))
    with client,pytest.raises(ProviderError,match=code):
        SmartRecruitersConnector().fetch(source(),PublicReader(client,RequestBudget(DiscoverySettings(retry_limit=0)),metrics()),{})


def test_normalization_dates_identity_and_unknown_payroll():
    connector=SmartRecruitersConnector();job,url,evidence=connector.normalize(source(),posting())
    assert job.location=='Noida, India' and job.work_mode=='hybrid'
    assert url.endswith('/Nagarro1/1') and evidence['description_pending']
    assert evidence['published_at']=='2026-10-15T01:00:00Z'
    assert 'direct_permanent_payroll' not in evidence
    with pytest.raises(ValueError,match='identity'):
        connector.normalize(source(),posting()|{'company':{'identifier':'Airtable'}})
    with pytest.raises(ValueError,match='country'):
        connector.normalize(source(),posting()|{'location':{'country':'us'}})
    with pytest.raises(ValueError,match='unsafe'):
        connector.normalize(source(),posting()|{'applyUrl':'https://evil.test/'})
    with pytest.raises(ValueError,match='generic'):
        connector.normalize(source(),posting()|{'applyUrl':'https://jobs.smartrecruiters.com/'})


@pytest.mark.parametrize('url',[
    'https://evil.test/?next=https://linkedin.com/jobs/view/123',
    'https://linkedin.com.evil.test/jobs/view/123',
    'https://linkedin.com/jobs/search/?keywords=react',
    'https://in.indeed.com/viewjob',
    'https://user:pass@linkedin.com/jobs/view/123',
    'https://linkedin.com:bad/jobs/view/123',
    'http://linkedin.com/jobs/view/123',
])
def test_alert_link_rejection(url):
    assert GmailExportAdapter._specific_url(url) is None


def test_provider_identity_and_exact_links():
    one=GmailExportAdapter._specific_url('https://in.indeed.com/viewjob?jk=abc123&utm_source=email')
    two=GmailExportAdapter._specific_url('https://in.indeed.com/viewjob?jk=def456')
    assert one.endswith('?jk=abc123') and canonical_url(one)!=canonical_url(two)
    assert GmailExportAdapter._specific_url('https://www.linkedin.com/jobs/view/react-lead-123456')
    assert GmailExportAdapter._specific_url('https://www.hirist.tech/j/publicis-sapient-senior-java-developer-1662376')
    assert GmailExportAdapter._specific_url('https://www.naukri.com/job-listings-react-lead-employer-noida-5-to-8-years-12345678901')


def amazon_post(identifier='123'):
    return {'id':'search-uuid','id_icims':identifier,'title':'Senior Frontend Engineer','company_name':'Amazon India entity',
        'country_code':'IND','location':'IN, Noida','job_path':f'/en/jobs/{identifier}/senior-frontend-engineer',
        'description':'Build UI','basic_qualifications':'React and TypeScript required','preferred_qualifications':'Python preferred',
        'posted_date':'October 15, 2026','job_schedule_type':'full-time','locations':['{"countryIso2a":"IN","city":"Noida","type":"ONSITE"}']}


def test_amazon_pagination_normalization_and_unknown_dates():
    def handler(req):
        offset=int(req.url.params['offset'])
        assert req.url.params['normalized_country_code[]']=='IND'
        return httpx.Response(200,json={'hits':2,'error':None,'jobs':[amazon_post(str(offset+1))]})
    connector=AmazonConnector()
    with httpx.Client(transport=httpx.MockTransport(handler)) as client:
        batch=connector.fetch(source(),PublicReader(client,RequestBudget(DiscoverySettings(page_size=1)),metrics()),{})
    assert batch.checkpoint['complete'] and len(batch.records)==2
    job,url,ev=connector.normalize(source(),batch.records[0])
    assert job.source_job_id=='1' and url.endswith('/en/jobs/1/senior-frontend-engineer')
    assert ev['published_at']=='2026-10-15' and ev['publication_precision']=='DAY'
    assert {'react','typescript'}<=set(ev['required']) and ev['preferred']==['python']
    raw=batch.records[0]|{'posted_date':'recently'}
    assert connector.normalize(source(),raw)[2]['published_at'] is None
    with pytest.raises(ValueError,match='country'):connector.normalize(source(),raw|{'country_code':'USA'})
    with pytest.raises(ValueError,match='path'):connector.normalize(source(),raw|{'job_path':'https://evil.test'})


def test_amazon_budget_cap_and_bad_pagination():
    connector=AmazonConnector();cfg=DiscoverySettings(max_requests=1,page_size=1,retry_limit=0)
    with httpx.Client(transport=httpx.MockTransport(lambda req:httpx.Response(200,json={'hits':3,'jobs':[amazon_post()]}))) as client:
        batch=connector.fetch(source(),PublicReader(client,RequestBudget(cfg),metrics()),{})
    assert batch.checkpoint['next_offset']==1 and not batch.checkpoint['complete']
    assert batch.checkpoint['termination_reason']=='request_budget_exhausted'
    with httpx.Client(transport=httpx.MockTransport(lambda req:httpx.Response(200,json={'hits':10000,'jobs':[amazon_post()]}))) as client:
        batch=connector.fetch(source(),PublicReader(client,RequestBudget(cfg),metrics()),{'next_offset':9999,'scope':{'country':'IND','query':'','sort':'recent'}})
    assert batch.checkpoint['termination_reason']=='search_index_result_cap' and not batch.checkpoint['complete']
    for payload in ({'error':'restricted','jobs':[]},{'hits':True,'jobs':[]},{'hits':1,'jobs':[]}):
        with httpx.Client(transport=httpx.MockTransport(lambda req:httpx.Response(200,json=payload))) as client,pytest.raises(ProviderError):
            connector.fetch(source(),PublicReader(client,RequestBudget(cfg),metrics()),{})
    with httpx.Client(transport=httpx.MockTransport(lambda req:httpx.Response(200,json={'hits':1,'jobs':[amazon_post('SF123')|{'id_icims':None}]}))) as client:
        batch=connector.fetch(source(),PublicReader(client,RequestBudget(cfg),metrics()),{})
    assert batch.records[0]['id']=='SF123'
    assert connector.normalize(source(),batch.records[0])[0].source_job_id=='SF123'


def test_smart_detail_cache_and_backlog():
    connector=SmartRecruitersConnector();row=posting();row.pop('releasedDate')
    detail=row|{'jobAd':{'sections':{'qualifications':{'title':'Qualifications','text':'React required.'}}}}
    with httpx.Client(transport=httpx.MockTransport(lambda req:httpx.Response(200,json=detail))) as client:
        batch=connector.enrich(source(),Batch([row],{}),PublicReader(client,RequestBudget(DiscoverySettings()),metrics()),{},Profile(locations=['Noida']))
    assert batch.checkpoint['detail_verified']==1
    job,url,evidence=connector.normalize(source(),row)
    assert evidence['required']==['react'] and not evidence['description_pending']
    cached={'1':{'job':job.model_dump(mode='json'),'evidence':evidence}}
    with httpx.Client(transport=httpx.MockTransport(lambda req:pytest.fail('cache should avoid network'))) as client:
        again=connector.enrich(source(),Batch([posting()|{'releasedDate':None}],{}),PublicReader(client,RequestBudget(DiscoverySettings()),metrics()),cached,Profile(locations=['Noida']))
    assert again.checkpoint['detail_cache_hits']==1
    assert connector.normalize(source(),again.records[0])[0].description==job.description
    cached['1']['evidence']['source_active']=False
    with httpx.Client(transport=httpx.MockTransport(lambda req:pytest.fail('cache should avoid network'))) as client:
        closed=connector.enrich(source(),Batch([posting()|{'releasedDate':None}],{}),PublicReader(client,RequestBudget(DiscoverySettings()),metrics()),cached,Profile(locations=['Noida']))
    assert connector.normalize(source(),closed.records[0])[2]['source_active'] is False
    with httpx.Client(transport=httpx.MockTransport(lambda req:httpx.Response(429,headers={'retry-after':'600'}))) as client:
        blocked=connector.enrich(source(),Batch([posting()|{'releasedDate':None}],{}),PublicReader(client,RequestBudget(DiscoverySettings(retry_limit=0)),metrics()),{},Profile(locations=['Noida']))
    assert blocked.checkpoint['pending_details']==[{'id':'1','reason':'rate_limited'}] and blocked.checkpoint['retry_after']==600


def test_multi_job_unquoted_plaintext_alert_and_report(tmp_path):
    db=Database(tmp_path/'state.sqlite3');engine=M2DEngine(db);engine.seed()
    message=EmailExportMessage(message_id='multi',received_at='2026-10-15T01:00:00Z',subject='React jobs',
        body='<a href=https://www.linkedin.com/jobs/view/react-123>React <b>lead</b></a>\nhttps://in.indeed.com/viewjob?jk=abc123\nhttps://in.indeed.com/viewjob?jk=def456\nhttps://linkedin.com/jobs/search/')
    first=GmailExportAdapter(engine).ingest([message]);second=GmailExportAdapter(engine).ingest([message])
    assert first['vacancies']==3 and second['vacancies']==0
    assert GmailExportAdapter(engine).ingest([message],reprocess=True)['vacancies']==0
    snapshot=report(db)
    assert snapshot['gmail']['observed_vacancy_hosts']['in.indeed.com']==2
    assert next(x for x in snapshot['families'] if x['id']=='linkedin')['state']=='NOT_IMPLEMENTED'
    assert next(x for x in snapshot['families'] if x['id']=='flexjobs')['state']=='PAID_ACCESS_EXCLUDED'
    assert not any(x['id']=='pipeline-handoff' for x in snapshot['families'])


def test_reextraction_journal_cursor_provider_and_idempotency(tmp_path):
    db=Database(tmp_path/'state.sqlite3');engine=M2DEngine(db);engine.seed()
    message=EmailExportMessage(message_id='existing',received_at='2026-10-15T01:00:00Z',sender_domain='alerts.naukri.com',subject='React lead',body='https://www.naukri.com/job-listings-react-lead-employer-12345678901')
    # Simulate old extractor journal without a vacancy observation.
    GmailExportAdapter(engine).ingest([message.model_copy(update={'body':'generic careers link'})])
    class Reader:
        def __init__(self):self.tokens=[]
        def read_interval(self,start,end,token,page_size):
            self.tokens.append(token)
            return {'messages':[message] if token is None else [],'next_page_token':'second' if token is None else None}
    reader=Reader();start='2026-10-15T00:00:00Z';end='2026-10-15T02:00:00Z'
    first=reextract_interval(engine,reader,start,end,max_pages=1)
    assert first['status']=='PARTIAL' and first['next_page_token']=='second' and first['new_observations']==1
    done=reextract_interval(engine,reader,start,end,max_pages=1)
    again=reextract_interval(engine,reader,start,end)
    assert done['status']=='COMPLETED' and again['reused'] and reader.tokens==[None,'second']
    assert done['alert_sources']['naukri.com']['messages']==1
    with db.transaction() as conn:
        assert conn.execute('select count(*) from gmail_intake').fetchone()[0]==1
        assert conn.execute("select count(*) from job_observations where stream='GMAIL'").fetchone()[0]==1
    for args in [(end,start),(start.replace('Z',''),end)]:
        with pytest.raises(ValueError):reextract_interval(engine,reader,*args)
    with pytest.raises(ValueError):reextract_interval(engine,reader,start,end,max_pages=0)


def test_official_detail_worker_lease_cache_and_preservation(tmp_path):
    db=Database(tmp_path/'state.sqlite3');cfg=DiscoveryConfig(sources=[source()],settings=DiscoverySettings(retry_limit=0))
    Registry(db).sync(cfg);service=JobService(db)
    with db.transaction() as conn:candidate=service._candidate(conn)
    raw=posting();raw.pop('releasedDate')
    persist_batch(db,source(),SmartRecruitersConnector(),[raw],candidate.search_background(),candidate.search_profile(),metrics(),RequestBudget(cfg.settings))
    def handler(req):
        return httpx.Response(200,json=raw|{'jobAd':{'sections':{'qualifications':{'title':'Qualifications','text':'React required.'}}}})
    discovery=Discovery(db,cfg.settings,lambda:httpx.Client(transport=httpx.MockTransport(handler)))
    try:
        first=discovery.execute_detail_checks(discovery.claim());again=discovery.execute_detail_checks(discovery.claim())
        assert first['status']=='COMPLETED' and first['metrics']['updated_jobs']==1
        assert again['metrics']['requests']==0 and again['metrics']['detail_cache_hits']==1
        snapshot=report(db)
        assert len(snapshot['employer_feeds'])==1 and snapshot['preserved_counts']['jobs']==1
    finally:discovery.close()


def test_release_local_journal_resume_export_and_no_drafts(tmp_path,monkeypatch):
    config=tmp_path/'sources.json';config.write_text(DiscoveryConfig().model_dump_json())
    args=Namespace(db=str(tmp_path/'state.sqlite3'),config=str(config),export_dir=str(tmp_path/'exports'),resume=None,gmail_live=False)
    first=release.execute(args)
    assert first['status']=='COMPLETED' and first['stages']['gmail_intake']['status']=='SKIPPED'
    assert first['stages']['postflight']['counts']['draft_operations']==0
    assert (tmp_path/'exports'/'source-coverage.json').is_file()
    assert list((tmp_path/'exports').glob('*.xlsx'))
    monkeypatch.setattr(Discovery,'execute',lambda *_:pytest.fail('completed stage must not replay'))
    args.resume=first['run_id'];again=release.execute(args)
    assert again['run_id']==first['run_id'] and again['status']=='COMPLETED'


def test_reextraction_recovers_failure_without_duplicate_observations(tmp_path):
    engine=M2DEngine(Database(tmp_path/'state.sqlite3'));engine.seed()
    class Reader:
        failed=False
        def read_interval(self,start,end,token,page_size):
            if not self.failed:self.failed=True;raise TimeoutError()
            return {'messages':[],'next_page_token':None}
    reader=Reader();start='2026-10-15T00:00:00Z';end='2026-10-15T01:00:00Z'
    with pytest.raises(TimeoutError):reextract_interval(engine,reader,start,end)
    assert reextract_interval(engine,reader,start,end)['status']=='COMPLETED'


def test_network_reserve_allows_metadata_persistence(tmp_path):
    clock=[0.0];settings=DiscoverySettings(max_runtime_seconds=10,retry_limit=0)
    budget=RequestBudget(settings,clock=lambda:clock[0]);clock[0]=9.1
    connector=SmartRecruitersConnector();raw=posting()|{'releasedDate':None}
    with httpx.Client(transport=httpx.MockTransport(lambda req:pytest.fail('reserved write time must not be spent on network'))) as client:
        batch=connector.enrich(source(),Batch([raw],{'complete':True}),PublicReader(client,budget,metrics()),{},Profile(locations=['Noida']))
    assert batch.checkpoint['pending_details'][0]['reason']=='runtime_budget_exhausted'
    db=Database(tmp_path/'state.sqlite3');service=JobService(db)
    with db.transaction() as conn:candidate=service._candidate(conn)
    persist_batch(db,source(),connector,batch.records,candidate.search_background(),candidate.search_profile(),metrics(),budget)
    with db.transaction() as conn:assert conn.execute('select count(*) from jobs').fetchone()[0]==1


def test_imported_text_is_not_an_excel_formula(tmp_path):
    from openpyxl import Workbook,load_workbook
    from job_agent.jobqueue import JobQueueWorkbook
    workbook=Workbook();JobQueueWorkbook._sheet(workbook,'Untrusted',('Title',),[{'Title':'=1+1'}])
    path=tmp_path/'review.xlsx';workbook.save(path)
    cell=load_workbook(path,data_only=False)['Untrusted']['A2']
    assert cell.value=='=1+1' and cell.data_type=='s'
