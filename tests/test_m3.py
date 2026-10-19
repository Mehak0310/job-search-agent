import json,sqlite3
from pathlib import Path
import pytest
from openpyxl import load_workbook
from job_agent.database import Database
from job_agent.m2d import M2DEngine
from job_agent.m2d_models import OperationalFacts,TriState
from job_agent.m3 import M3Engine
from job_agent.m3_export import M3Workbook
from job_agent.m3_models import ContactInput,DraftRequest,EmployerQuestion,PacketEdit
from job_agent.service import JobService
from job_agent.models import JobInput
from job_agent.gmail_local import GmailDraftAdapter,DRAFT_SCOPE

def setup_engine(tmp_path,readiness="READY"):
    db=Database(tmp_path/"state.sqlite3");service=JobService(db);__import__('tests.synthetic_profile',fromlist=['seed']).seed(db)
    service.import_rows([JobInput(source="fixture",source_job_id="role-1",company="Example Corp",title="Senior Frontend Engineer",location="Noida",work_mode="hybrid",description="React TypeScript leadership",skills=["react","typescript","graphql"],application_url="https://example.com/jobs/role-1",min_experience=5,max_experience=8,is_sample=True)])
    with db.transaction() as conn:
        row=conn.execute("SELECT config FROM candidate_profiles_v2 WHERE id=1").fetchone();profile=json.loads(row[0]);profile["confirmed_facts"]={"notice_period":"45 days; not serving","expected_compensation":"42 LPA; not a hard minimum"};conn.execute("UPDATE candidate_profiles_v2 SET config=? WHERE id=1",(json.dumps(profile),))
    m2d=M2DEngine(db);m2d.seed();m2d.sync_existing_jobs()
    with db.transaction() as conn:key=conn.execute("SELECT job_key FROM job_operational_state").fetchone()[0]
    facts=OperationalFacts(active=TriState.PASS,india_eligible=TriState.UNKNOWN,headcount=TriState.UNKNOWN,direct_permanent_payroll=TriState.UNKNOWN,mandatory_requirements=TriState.UNKNOWN)
    if readiness=="EXCLUDED":facts=facts.model_copy(update={"india_eligible":TriState.FAIL})
    m2d.set_operational_state(key,facts,job_id=key,strong_fit=True)
    docs=tmp_path/"documents"/"manifest.json";docs.parent.mkdir();docs.write_text(json.dumps({"documents":[{"canonical_name":"Resume.pdf","version":"v3","sha256":"synthetic","status":"approved recruiter resume"},{"canonical_name":"Profile_Card.png","version":"current","sha256":"synthetic","status":"approved reference"},{"canonical_name":"Answer_Bank.docx","version":"v6","status":"missing private answer bank"}]}))
    return db,M3Engine(db,docs),key

def verified_contact(engine,key):
    return engine.add_contact(key,ContactInput(name="Ria Recruiter",employer="Example Corp",email="ria@example.com",source_url="https://example.com/team/ria",role_relevance="Talent acquisition for Example Corp engineering",verification_status="VERIFIED"))


def test_packet_history_retains_submitted_jobs_without_reopening_active_queue(tmp_path):
    db,engine,key=setup_engine(tmp_path)
    packet=engine.prepare(key)
    with db.transaction() as conn:
        conn.execute("UPDATE jobs SET status='SUBMITTED' WHERE id=?",(packet['job_id'],))
        before=conn.execute('SELECT COUNT(*) FROM status_history').fetchone()[0]
    assert not engine.list()
    assert engine.list(include_history=True)[0]['packet_id']==packet['packet_id']
    with db.transaction() as conn:
        assert conn.execute('SELECT status FROM jobs WHERE id=?',(packet['job_id'],)).fetchone()[0]=='SUBMITTED'
        assert conn.execute('SELECT COUNT(*) FROM status_history').fetchone()[0]==before


def test_inspected_portal_packet_uses_actual_documents_and_preserves_edits(tmp_path):
    db,engine,key=setup_engine(tmp_path)
    form={'required_documents':['Resume.pdf'],'missing_personal_answers':['employment restrictions'],'outstanding_agent_checks':['Permanent payroll unknown']}
    M2DEngine(db).record_evidence(key,'portal_form','PASS',factual=json.dumps(form),evidence_url='https://example.com/jobs/role-1')
    packet=engine.prepare(key)
    assert packet['required_documents']==['Resume.pdf']
    assert 'employment restrictions' in packet['missing_information']
    assert 'Permanent payroll unknown' in packet['version_data']['gaps']
    assert 'recipient' not in packet['manual_next_action']
    value=PacketEdit(subject='My subject',email_text='My text',email_html='<p>My text</p>',supporting_statement='My statement')
    edited=engine.edit(packet['packet_id'],value)
    assert engine.prepare(key)['current_version']==edited['current_version']
    assert engine.get(packet['packet_id'])['version_data']['supporting_statement']=='My statement'

def test_packet_grounding_instructions_and_no_regeneration(tmp_path):
    db,engine,key=setup_engine(tmp_path);packet=engine.prepare(key)
    assert packet["status"]=="NEEDS_USER_INFORMATION" and "private_answer_bank_path" in packet["missing_information"]
    version=packet["version_data"];assert "React" in version["email_text"] or "react" in version["email_text"]
    assert "<profile card here>" in version["email_text"] and "ATTACH: Resume.pdf." in version["email_text"]
    assert 'color:#c00' in version["email_html"] and "portfolio" not in version["email_text"].lower()
    assert engine.prepare(key)["current_version"]==1
    before=db.path.stat().st_size;result=engine.prepare_queue();assert result["selected"]==0 and result["unchanged"]==0 and db.path.stat().st_size>=before  # UNKNOWN stays in verification under the R2 addendum.

def test_known_incompatible_is_blocked(tmp_path,monkeypatch):
    monkeypatch.setattr('job_agent.runtime_policy.require_external_write',lambda operation:None)
    _,engine,key=setup_engine(tmp_path,"EXCLUDED");packet=engine.prepare(key)
    assert packet["status"]=="BLOCKED_INCOMPATIBLE"
    with pytest.raises(ValueError,match="verified_recipient_required"):engine.create_draft(packet["packet_id"],DraftRequest(contact_id=engine.add_contact(key,ContactInput(name="No Reply",employer="Example Corp",email="noreply@example.com",source_url="https://example.com",role_relevance="automated sender",verification_status="NO_REPLY"))),object())

def test_questions_require_real_provenance_and_enforce_limits(tmp_path):
    _,engine,key=setup_engine(tmp_path);packet=engine.prepare(key)
    qid=engine.add_question(packet["packet_id"],EmployerQuestion(question="What is your notice period?",source="employer_portal",source_ref="https://example.com/jobs/role-1#notice",character_limit=100,field_kind="NOTICE"))
    refreshed=engine.get(packet["packet_id"]);q=next(x for x in refreshed["questions"] if x["question_id"]==qid)
    assert q["answer_status"]=="PROPOSED" and "45 days" in q["answer"] and q["source_ref"].startswith("https://")
    unknown=engine.add_question(packet["packet_id"],EmployerQuestion(question="Are you authorised to work here?",source="employer_portal",source_ref="field:work-auth",character_limit=3,field_kind="WORK_AUTHORIZATION"))
    assert next(x for x in engine.get(packet["packet_id"])["questions"] if x["question_id"]==unknown)["answer_status"]=="NEEDS_USER_INFORMATION"
    with pytest.raises(ValueError,match="answer_exceeds_character_limit"):engine.answer_question(unknown,"yes, permanently")

def test_contact_verification_draft_success_duplicate_edit_preservation(tmp_path,monkeypatch):
    # Isolated adapter simulation of the historical journal, never live permission.
    monkeypatch.setattr('job_agent.runtime_policy.require_external_write',lambda operation:None)
    _,engine,key=setup_engine(tmp_path);packet=engine.prepare(key)
    with pytest.raises(ValueError,match="verified_contact_requires"):engine.add_contact(key,ContactInput(name="Guess",employer="Example Corp",email="guess@example.com",role_relevance="unknown",verification_status="VERIFIED"))
    contact=verified_contact(engine,key);assert engine.get(packet["packet_id"])["status"]=="REVIEW_DRAFT_READY"
    class Adapter:
        def __init__(self):self.calls=[]
        def create_draft(self,**kwargs):self.calls.append(kwargs);return {"id":"draft-synthetic-1","message":{"id":"message-synthetic-1"}}
    adapter=Adapter();saved=engine.create_draft(packet["packet_id"],DraftRequest(contact_id=contact),adapter)
    assert saved["status"]=="GMAIL_DRAFT_SAVED" and saved["manual_files_pending"] and len(adapter.calls)==1
    with pytest.raises(ValueError,match="duplicate_draft_operation"):engine.create_draft(packet["packet_id"],DraftRequest(contact_id=contact),adapter)
    current=engine.get(packet["packet_id"]);v=current["version_data"]
    edited=engine.edit(packet["packet_id"],PacketEdit(subject=v["subject"]+" edited",email_text=v["email_text"],email_html=v["email_html"],supporting_statement=v["supporting_statement"]))
    assert edited["current_version"]==2 and edited["draft_operations"][0]["packet_version"]==1 and edited["draft_operations"][0]["state"]=="SAVED"

def test_verified_contact_rejects_automated_or_unrelated_addresses(tmp_path):
    _,engine,key=setup_engine(tmp_path);engine.prepare(key)
    with pytest.raises(ValueError,match="role_related_person"):
        engine.add_contact(key,ContactInput(name="Notifications",employer="Example Corp",email="no-reply@example.com",source_url="https://example.com",role_relevance="Recruiting notification mailbox",verification_status="VERIFIED"))
    with pytest.raises(ValueError,match="role_relevance_required"):
        engine.add_contact(key,ContactInput(name="Engineer",employer="Example Corp",email="person@example.com",source_url="https://example.com",role_relevance="Software engineer",verification_status="VERIFIED"))

def test_packet_regeneration_preserves_current_user_edit(tmp_path):
    _,engine,key=setup_engine(tmp_path);packet=engine.prepare(key);v=packet['version_data']
    edited=engine.edit(packet['packet_id'],PacketEdit(subject='My reviewed subject',email_text=v['email_text'],email_html=v['email_html'],supporting_statement=v['supporting_statement']))
    assert engine.prepare(key)['current_version']==edited['current_version']
    assert engine.get(packet['packet_id'])['version_data']['subject']=='My reviewed subject'

def test_stored_contact_research_accepts_hiring_address_and_rejects_accessibility(tmp_path):
    db,engine,key=setup_engine(tmp_path);engine.prepare(key)
    with db.transaction() as conn:
        conn.execute("UPDATE jobs SET description=? WHERE id=?",("For application or interview assistance contact careers@example.com. For a reasonable accommodation contact candidate_accessibility@example.com.",key))
    result=engine.research_stored_contacts();packet=engine.by_job(key)
    assert result['verified_new']==1 and result['rejected']==1
    assert [contact['email'] for contact in packet['contacts']]==['careers@example.com']

def test_uncertain_outcome_requires_manual_reconciliation(tmp_path,monkeypatch):
    monkeypatch.setattr('job_agent.runtime_policy.require_external_write',lambda operation:None)
    _,engine,key=setup_engine(tmp_path);packet=engine.prepare(key);contact=verified_contact(engine,key)
    class Timeout:
        def create_draft(self,**kwargs):raise TimeoutError
    with pytest.raises(ValueError,match="outcome_uncertain"):engine.create_draft(packet["packet_id"],DraftRequest(contact_id=contact),Timeout())
    op=engine.get(packet["packet_id"])["draft_operations"][0];assert op["state"]=="UNCERTAIN"
    assert engine.resolve_uncertain(op["operation_id"],"NO_DRAFT")["state"]=="RESOLVED_NO_DRAFT"

def test_export_and_schema_preserve_existing_records(tmp_path):
    db,engine,key=setup_engine(tmp_path);engine.prepare(key);path=M3Workbook(engine).export(tmp_path/"exports")["path"]
    book=load_workbook(path,read_only=True);assert book.sheetnames==["Application Queue","Application Text","Questions Answers","Contacts Evidence","Missing Information","Run Audit"]
    assert book["Application Queue"]["A2"].value==key
    with db.transaction() as conn:
        assert conn.execute("PRAGMA user_version").fetchone()[0]==17
        assert conn.execute("PRAGMA integrity_check").fetchone()[0]=="ok"
        assert conn.execute("SELECT count(*) FROM jobs").fetchone()[0]==1

def test_gmail_adapter_boundary_has_no_send_and_requires_local_consent(tmp_path):
    adapter=GmailDraftAdapter(tmp_path/"client.json",tmp_path/"token.json")
    assert adapter.status()["scope"]==DRAFT_SCOPE and "sending" in adapter.status()["warning"] and not hasattr(adapter,"send")
    with pytest.raises(ValueError,match="external_writes_disabled"):adapter.create_draft(to="r@example.com",subject="s",text="t",html="<p>t</p>")
