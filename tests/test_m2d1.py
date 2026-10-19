import base64,zipfile
from pathlib import Path
from docx import Document
from job_agent.candidate import CandidateProfile
from job_agent.document_manifest import resume_safety
from job_agent.gmail_local import READONLY_SCOPE,DRAFT_SCOPE,GmailOAuthReader,gmail_query,message_from_api
import pytest
from job_agent.public_connectors import GreenhouseConnector
from job_agent.resume_review import extract
from job_agent.discovery_models import SourceConfig

def encoded(value):return base64.urlsafe_b64encode(value.encode()).decode().rstrip('=')

def test_gmail_read_scope_query_received_only_and_cursor():
    query=gmail_query('2026-10-14T00:00:00Z','2026-10-15T00:00:00Z')
    assert '-in:sent -in:drafts' in query and READONLY_SCOPE.endswith('gmail.readonly') and DRAFT_SCOPE not in query
    raw={'id':'m1','threadId':'t1','internalDate':'1791936000000','labelIds':['INBOX'],'payload':{'headers':[{'name':'Subject','value':'Roles'}],'mimeType':'text/plain','body':{'data':encoded('Job: Lead\nCompany: Example\nURL: https://example.com/jobs/1')}}}
    message=message_from_api(raw,'next');assert message and message.message_id=='m1' and message.cursor=='next' and message.mailbox=='received'
    assert message_from_api(raw|{'labelIds':['SENT']}) is None and message_from_api(raw|{'labelIds':['DRAFT']}) is None

def test_gmail_oauth_missing_local_configuration_is_explicit(tmp_path):
    reader=GmailOAuthReader(tmp_path/'client.json',tmp_path/'token.json');assert reader.status()['live_ready'] is False
    with pytest.raises(ValueError,match='gmail_client_secret_missing'):reader.authorize()
    with pytest.raises(ValueError,match='gmail_local_authorization_required'):reader.read_interval('2026-10-14T00:00:00Z','2026-10-15T00:00:00Z')

def test_effective_profile_preserves_evidence_categories_and_durations():
    profile=CandidateProfile(total_experience_years=8,technology_years={'LLM':3,'React':5},ai_employment_years=None,backend_experience_years=3,search_skills=['graphql'])
    search=profile.search_profile();background=profile.search_background()
    assert search.experience_years==8 and 'graphql' in search.skills and profile.technology_years=={'llm':3,'react':5}
    assert background.backend_years==3 and profile.ai_employment_years is None

def test_docx_extraction_includes_headers_and_drawing_text(tmp_path):
    path=tmp_path/'resume.docx';doc=Document();doc.add_paragraph('Angular experience');doc.sections[0].header.paragraphs[0].text='US onsite assignment';doc.save(path)
    rebuilt=tmp_path/'rebuilt.docx'
    with zipfile.ZipFile(path) as source,zipfile.ZipFile(rebuilt,'w') as target:
        for item in source.infolist():
            value=source.read(item.filename)
            if item.filename=='word/document.xml':value=value.decode().replace('</w:body>','<w:p><w:r><w:t>Text box portfolio-free note</w:t></w:r></w:p></w:body>').encode()
            target.writestr(item,value)
    rebuilt.replace(path)
    values=' '.join(value for _,value in extract(path));assert 'US onsite assignment' in values and 'Text box portfolio-free note' in values

def test_greenhouse_long_location_is_compacted_but_preserved_as_evidence(greenhouse_record):
    raw='; '.join(f'Office {i}' for i in range(80));record=greenhouse_record|{'location':{'name':raw}}
    job,_,evidence=GreenhouseConnector().normalize(SourceConfig(id='gh',provider='greenhouse',company='Example',board='example'),record)
    assert len(job.location)<=200 and job.location.endswith('...') and evidence['source_location']==raw[:2000]
