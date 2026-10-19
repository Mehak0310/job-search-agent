import io,sqlite3
from docx import Document
from pypdf import PdfWriter
from job_agent.candidate import CandidateProfile
from job_agent.evaluation import DATASET,evaluate
from job_agent.intelligence import SentenceTransformerProvider

def docx_bytes():
    document=Document();document.add_paragraph('Independent Project: RAG application with FastAPI and Python');document.add_paragraph('Master degree in Data Science');buffer=io.BytesIO();document.save(buffer);return buffer.getvalue()

def pdf_bytes():
    writer=PdfWriter();writer.add_blank_page(width=612,height=792);buffer=io.BytesIO();writer.write(buffer);return buffer.getvalue()

def test_explicit_resume_parse_review_and_version(client):
    payload=docx_bytes();attached=client.post('/api/candidate-profile/resume',files={'file':('resume.docx',payload)});assert attached.status_code==200
    assert client.get('/api/candidate-profile/facts').json()==[]
    parsed=client.post('/api/candidate-profile/resume/parse');assert parsed.status_code==200 and parsed.json()['review_required']
    facts=client.get('/api/candidate-profile/facts').json();assert facts and all(x['status']=='PENDING' for x in facts)
    assert all(x['extraction_id']==parsed.json()['extraction_id'] and (' block ' in x['source_ref']) for x in facts)
    profile_before=client.get('/api/candidate-profile').json();skill=next(x for x in facts if x['category']=='SKILL_UNVERIFIED')
    assert client.patch('/api/candidate-profile/facts/'+skill['id'],json={'action':'ACCEPTED','reviewed_value':'  '}).status_code==422
    assert client.patch('/api/candidate-profile/facts/'+skill['id'],json={'action':'ACCEPTED','reviewed_value':skill['value']}).json()['status']=='ACCEPTED'
    profile_after=client.get('/api/candidate-profile').json();assert profile_after['version']==profile_before['version']+1 and skill['value'] in profile_after['project_skills']
    education=next(x for x in facts if x['category']=='EDUCATION');assert client.patch('/api/candidate-profile/facts/'+education['id'],json={'action':'REJECTED'}).status_code==200
    for fact in client.get('/api/candidate-profile/facts').json():
        if fact['status']=='PENDING':assert client.patch('/api/candidate-profile/facts/'+fact['id'],json={'action':'REJECTED'}).status_code==200
    assert client.get('/api/candidate-profile').json()['resume_reviewed'] is True
    assert client.patch('/api/candidate-profile/facts/'+education['id'],json={'action':'ACCEPTED'}).status_code==404

def test_pdf_parse_is_explicit_repeatable_and_provenanced(client):
    attached=client.post('/api/candidate-profile/resume',files={'file':('resume.pdf',pdf_bytes())});assert attached.status_code==200
    assert client.get('/api/candidate-profile/facts').json()==[]
    first=client.post('/api/candidate-profile/resume/parse');second=client.post('/api/candidate-profile/resume/parse')
    assert first.status_code==200 and second.status_code==200
    assert first.json()['extraction_id']==second.json()['extraction_id'] and first.json()['proposals']==0

def test_resume_parse_requires_attachment(client):
    response=client.post('/api/candidate-profile/resume/parse');assert response.status_code==422 and response.json()['detail']=='resume_not_attached'

def test_v6_migration_backup_and_audit(app):
    db=app.state.service.db
    with db.transaction() as conn:assert conn.execute('pragma user_version').fetchone()[0]==17
    assert db.path.with_suffix(db.path.suffix+'.pre-v6.bak').exists()

def test_trained_provider_is_optional_and_offline():
    try:SentenceTransformerProvider(model_id='missing-local-model',local_files_only=True)
    except (ImportError,ModuleNotFoundError,OSError):pass

def test_evaluation_accepts_provider_and_keeps_held_out_split():
    class ConstantProvider:
        model_id='test';version='1'
        def embed_batch(self,texts):return [[1.0,0.0] for _ in texts]
    result=evaluate(ConstantProvider())
    assert len(DATASET)==60 and result['split']=={'tuning':24,'held_out':36}
    assert result['held_out']['jobs_evaluated']==36
    assert result['held_out']['section']['hard_eligibility_violations']==0
