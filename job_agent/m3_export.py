"""Reviewable local M3 workbook; export is never proof of upload or submission."""
import hashlib,json
from datetime import datetime,timedelta,timezone
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font
from .database import now

class M3Workbook:
    def __init__(self,engine):self.engine=engine
    @staticmethod
    def _sheet(book,name,headers,rows):
        sheet=book.create_sheet(name);sheet.append(headers)
        for cell in sheet[1]:cell.font=Font(bold=True)
        for row in rows:sheet.append([row.get(h) for h in headers])
        sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions;return sheet
    def export(self,out_dir:Path):
        packets=self.engine.list();out_dir.mkdir(parents=True,exist_ok=True);stamp=datetime.now(timezone(timedelta(hours=5,minutes=30))).strftime("%Y-%m-%d_%H%M_IST")
        path=out_dir/f"Application_Review_{stamp}.xlsx";book=Workbook();book.remove(book.active)
        jobs=[];texts=[];questions=[];contacts=[];missing=[];audit=[]
        for summary in packets:
            packet=self.engine.get(summary["packet_id"]);v=packet["version_data"];url=packet["canonical_application_url"] or packet["application_url"]
            jobs.append({"JobKey":packet["job_key"],"Company":packet["company"],"Role":packet["title"],"ExactVacancyURL":url,"Status":packet["status"],"PacketVersion":packet["current_version"],"ProfileVersion":packet["profile_version"],"VacancyVersion":packet["vacancy_version"]})
            texts.append({"JobKey":packet["job_key"],"Subject":v["subject"],"EmailText":v["email_text"],"SupportingStatement":v["supporting_statement"],"Generator":v["generator"],"UserEdited":v["user_edited"],"Provenance":json.dumps(v["provenance"],ensure_ascii=False,sort_keys=True)})
            for q in packet["questions"]:questions.append({"JobKey":packet["job_key"],"Question":q["question"],"Source":q["source"],"SourceRef":q["source_ref"],"CharacterLimit":q["character_limit"],"Answer":q["answer"],"AnswerStatus":q["answer_status"]})
            for c in packet["contacts"]:contacts.append({"JobKey":packet["job_key"],"Name":c["name"],"Employer":c["employer"],"Email":c["email"],"RoleRelevance":c["role_relevance"],"Verification":c["verification_status"],"SourceURL":c["source_url"],"SourceMessageID":c["source_message_id"],"VerifiedAt":c["verified_at"]})
            for item in packet["missing_information"]:missing.append({"JobKey":packet["job_key"],"MissingInformation":item,"NextAction":packet["manual_next_action"]})
            audit.append({"JobKey":packet["job_key"],"PacketID":packet["packet_id"],"Version":packet["current_version"],"CreatedAt":packet["created_at"],"UpdatedAt":packet["updated_at"],"DraftOperations":len(packet["draft_operations"]),"ExportedAt":now(),"UploadStatus":"NOT_UPLOADED"})
        job_headers=["JobKey","Company","Role","ExactVacancyURL","Status","PacketVersion","ProfileVersion","VacancyVersion"]
        sheet=self._sheet(book,"Application Queue",job_headers,jobs)
        for i,row in enumerate(jobs,2):
            if row["ExactVacancyURL"]:sheet.cell(i,4).hyperlink=row["ExactVacancyURL"];sheet.cell(i,4).style="Hyperlink"
        self._sheet(book,"Application Text",["JobKey","Subject","EmailText","SupportingStatement","Generator","UserEdited","Provenance"],texts)
        self._sheet(book,"Questions Answers",["JobKey","Question","Source","SourceRef","CharacterLimit","Answer","AnswerStatus"],questions)
        self._sheet(book,"Contacts Evidence",["JobKey","Name","Employer","Email","RoleRelevance","Verification","SourceURL","SourceMessageID","VerifiedAt"],contacts)
        self._sheet(book,"Missing Information",["JobKey","MissingInformation","NextAction"],missing)
        self._sheet(book,"Run Audit",["JobKey","PacketID","Version","CreatedAt","UpdatedAt","DraftOperations","ExportedAt","UploadStatus"],audit)
        book.save(path);checksum=hashlib.sha256(path.read_bytes()).hexdigest()
        manifest=out_dir/"manual-transfer-manifest.json";data=json.loads(manifest.read_text("utf-8")) if manifest.exists() else {"schema":"JobQueue-v1","uploads":[],"exports":[]};data["exports"].append({"filename":path.name,"sha256":checksum,"stream":"M3_APPLICATION_REVIEW","created_at":now(),"upload_status":"NOT_UPLOADED"});manifest.write_text(json.dumps(data,indent=2)+"\n","utf-8")
        return {"path":path,"checksum":checksum,"packets":len(packets),"upload_status":"NOT_UPLOADED"}
