"""JobQueue-v1 local XLSX handoff. Workbooks are artifacts, never upload proof."""
from __future__ import annotations
import hashlib,json
from datetime import datetime,timedelta,timezone
from pathlib import Path
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font
from .database import now
from .m2d import M2DEngine,SNAPSHOT_SCHEMA
from .m2d_models import JobQueueImportRow

JOBS_COLUMNS=("JobKey","Lifecycle","Eligibility","Verification","Readiness","Company","Title","Location","WorkMode","ExactVacancyURL","ProviderIDs","Streams","SourceProvenance","PublishedAt","PublishedPrecision","EmailReceivedAt","FirstDiscoveredAt","LastSeenAt","LastChangedAt","FitScore","FitEvidence","Gaps","EligibilityEvidence","MissingFacts","NextAction")
COMMON_SHEETS=("Jobs","Draft Inputs","Verification Queue","Source Coverage","Run Audit")
PREFIX={"MAJOR":"Public_Jobs","GMAIL":"Gmail_Jobs","SPECIALIST":"Specialist_Jobs","REMOTE":"Remote_Jobs","CONSOLIDATED":"Consolidated_Review"}

def _json(value):return json.dumps(value,ensure_ascii=False,sort_keys=True,separators=(",",":"))
def _checksum(path):return hashlib.sha256(path.read_bytes()).hexdigest()

class JobQueueWorkbook:
    def __init__(self,engine:M2DEngine):self.engine=engine

    def _rows(self,stream):
        with self.engine.db.transaction() as conn:
            states=conn.execute("""SELECT s.*,j.company,j.title,j.location,j.work_mode,j.canonical_application_url,j.application_url,j.published_at,j.timestamp_confidence,j.first_seen,j.last_seen,j.last_changed,j.match_score,j.match_detail,j.eligibility_evidence
                FROM job_operational_state s LEFT JOIN jobs j ON j.id=s.job_id ORDER BY s.readiness,j.match_score DESC,s.job_key""").fetchall()
            observations=(conn.execute("SELECT * FROM job_observations ORDER BY job_key,source_id,observation_id").fetchall() if stream=="CONSOLIDATED" else conn.execute("SELECT * FROM job_observations WHERE stream=? ORDER BY job_key,source_id,observation_id",(stream,)).fetchall())
        by_key={}
        for row in observations:by_key.setdefault(row["job_key"],[]).append(row)
        output=[]
        for state in states:
            obs=by_key.get(state["job_key"],[])
            if not obs:continue
            evidence=[json.loads(row["evidence"]) for row in obs];first=evidence[0] if evidence else {}
            url=state["canonical_application_url"] or state["application_url"] or next((row["canonical_url"] for row in obs if row["canonical_url"]),None)
            missing=json.loads(state["missing_facts"]);readiness=state["readiness"]
            output.append(dict(zip(JOBS_COLUMNS,(state["job_key"],state["lifecycle"],state["eligibility"],state["verification"],readiness,state["company"] or first.get("company"),state["title"] or first.get("title"),state["location"] or first.get("location"),state["work_mode"] or "unknown",url,_json([row["provider_id"] for row in obs if row["provider_id"]]),_json(sorted({row["stream"] for row in obs})),_json([{"source":row["source_id"],"observed_at":row["observed_at"]} for row in obs]),state["published_at"] or first.get("published_at"),state["timestamp_confidence"] or first.get("publication_precision"),next((row["email_received_at"] for row in obs if row["email_received_at"]),None),state["first_seen"] or min(row["observed_at"] for row in obs),state["last_seen"] or max(row["observed_at"] for row in obs),state["last_changed"],state["match_score"],state["match_detail"] or "{}",state["exclusion_reason"] or "",state["eligibility_evidence"] or "",_json(missing),"Verify: "+", ".join(missing) if missing else "Manual portal review"))))
        return output

    @staticmethod
    def _sheet(workbook,name,headers,rows):
        sheet=workbook.create_sheet(name);sheet.append(list(headers))
        for cell in sheet[1]:cell.font=Font(bold=True)
        for row in rows:
            values=[row.get(header) for header in headers];sheet.append(values)
            for cell,value in zip(sheet[sheet.max_row],values):
                if isinstance(value,str):cell.data_type='s'
        sheet.freeze_panes="A2";sheet.auto_filter.ref=sheet.dimensions
        return sheet

    def export(self,stream,out_dir:Path,run_id="manual"):
        if stream not in PREFIX:raise ValueError("unsupported_stream")
        rows=self._rows(stream);out_dir.mkdir(parents=True,exist_ok=True);stamp=datetime.now(timezone(timedelta(hours=5,minutes=30))).strftime("%Y-%m-%d_%H%M_IST");path=out_dir/f"{PREFIX[stream]}_{stamp}.xlsx"
        workbook=Workbook();workbook.remove(workbook.active)
        visible=[row for row in rows if row["Readiness"]!="EXCLUDED"];jobs=self._sheet(workbook,"Jobs",JOBS_COLUMNS,visible)
        link_col=JOBS_COLUMNS.index("ExactVacancyURL")+1
        for index,row in enumerate(visible,2):
            if row["ExactVacancyURL"]:jobs.cell(index,link_col).hyperlink=row["ExactVacancyURL"];jobs.cell(index,link_col).style="Hyperlink"
        draft_headers=("JobKey","Company","Title","ExactVacancyURL","VerifiedFacts","MissingFacts","Constraints")
        self._sheet(workbook,"Draft Inputs",draft_headers,[{"JobKey":r["JobKey"],"Company":r["Company"],"Title":r["Title"],"ExactVacancyURL":r["ExactVacancyURL"],"VerifiedFacts":_json({"lifecycle":r["Lifecycle"],"eligibility":r["Eligibility"],"verification":r["Verification"]}),"MissingFacts":r["MissingFacts"],"Constraints":"No M3 prose; no submission authorization"} for r in visible])
        tasks=self.engine.tasks();task_headers=("TaskID","Kind","JobKey","SourceID","Stream","Category","Status","Attempts","DueAt","Blocker","Checkpoint")
        self._sheet(workbook,"Verification Queue",task_headers,[{"TaskID":r["task_id"],"Kind":r["kind"],"JobKey":r["job_key"],"SourceID":r["source_id"],"Stream":r["stream"],"Category":r["category"],"Status":r["status"],"Attempts":r["attempts"],"DueAt":r["due_at"],"Blocker":r["blocker"],"Checkpoint":r["checkpoint"]} for r in tasks if r["kind"]=="VERIFICATION"])
        coverage_headers=("SourceID","DisplayName","Stream","Owner","ImplementationState","CoverageStatus","LastAttempt","MeaningfulSuccess","NextDue","Blocker","Backlog","BoundedScope","Cursor")
        self._sheet(workbook,"Source Coverage",coverage_headers,[{k:v for k,v in zip(coverage_headers,(r["source_id"],r["display_name"],r["stream"],r["owner"],r["implementation_state"],r["coverage_status"],r["last_attempt"],r["meaningful_success"],r["next_due"],r["blocker"],r["backlog"],r["bounded_scope"],r["cursor"]))} for r in self.engine.sources() if stream=="CONSOLIDATED" or r["stream"]==stream or r["owner"]==stream])
        with self.engine.db.transaction() as conn:snapshots=[dict(row) for row in conn.execute("SELECT * FROM m2d_snapshots WHERE stream=? ORDER BY created_at",(stream,))]
        audit_headers=("SnapshotID","RunID","SchemaVersion","CompletionStatus","StartedAt","FinishedAt","DistinctJobs","Observations","Counts","ArtifactName","ArtifactChecksum")
        self._sheet(workbook,"Run Audit",audit_headers,[{k:v for k,v in zip(audit_headers,(r["snapshot_id"],r["run_id"],r["schema_version"],r["completion_status"],r["started_at"],r["finished_at"],r["distinct_jobs"],r["observations"],r["counts"],r["artifact_name"],r["artifact_checksum"]))} for r in snapshots])
        if stream in {"MAJOR","CONSOLIDATED"}:
            employer_headers=("EmployerKey","Name","CareersURL","Headcount","HeadcountState","EvidenceURL","VerifiedAt")
            self._sheet(workbook,"Employer Careers Registry",employer_headers,[{k:v for k,v in zip(employer_headers,(r["employer_key"],r["name"],r["careers_url"],r["headcount"],r["headcount_state"],r["evidence_url"],r["verified_at"]))} for r in self.engine.employers()])
            shortlist=[r for r in visible if r["Readiness"] in {"FULLY_VERIFIED_APPLICATION_READY","READY_FOR_PORTAL_REVIEW"}]
            self._sheet(workbook,"Application Review Shortlist",JOBS_COLUMNS,shortlist)
        excluded=[r for r in rows if r["Readiness"]=="EXCLUDED"]
        if excluded:self._sheet(workbook,"Excluded",JOBS_COLUMNS,excluded)
        workbook.save(path);checksum=_checksum(path)
        manifest=out_dir/"manual-transfer-manifest.json";data=json.loads(manifest.read_text("utf-8")) if manifest.exists() else {"schema":SNAPSHOT_SCHEMA,"uploads":[],"exports":[]};data["exports"].append({"filename":path.name,"sha256":checksum,"stream":stream,"created_at":now(),"upload_status":"NOT_UPLOADED"});manifest.write_text(json.dumps(data,indent=2)+"\n",encoding="utf-8")
        snapshot=self.engine.snapshot(run_id=run_id,stream=stream,status="PARTIAL",started_at=now(),counts={"jobs":len(visible),"excluded":len(excluded)},artifact_name=path.name,artifact_checksum=checksum)
        return {"path":path,"checksum":checksum,"snapshot_id":snapshot,"jobs":len(visible),"excluded":len(excluded),"upload_status":"NOT_UPLOADED"}

    def import_workbook(self,path:Path,stream,report_identity=None):
        if path.stat().st_size>20*1024*1024:raise ValueError("workbook_too_large")
        checksum=_checksum(path)
        with self.engine.db.transaction() as conn:
            existing=conn.execute("SELECT * FROM m2d_imports WHERE checksum=?",(checksum,)).fetchone()
            if existing:return {"import_id":existing["import_id"],"rows_imported":existing["rows_imported"],"duplicate":True}
        workbook=load_workbook(path,read_only=True,data_only=True)
        missing=[name for name in COMMON_SHEETS if name not in workbook.sheetnames]
        if missing:raise ValueError("missing_sheets:"+",".join(missing))
        sheet=workbook["Jobs"];headers=tuple(cell.value for cell in next(sheet.iter_rows(min_row=1,max_row=1)))
        if not set(JOBS_COLUMNS)<=set(headers):raise ValueError("invalid_jobqueue_columns")
        rows=[]
        for values in sheet.iter_rows(min_row=2,values_only=True):
            raw=dict(zip(headers,values));rows.append(JobQueueImportRow(job_key=str(raw["JobKey"]),company=str(raw["Company"]),title=str(raw["Title"]),exact_vacancy_url=str(raw["ExactVacancyURL"]),imported_status_claim=str(raw["Readiness"]) if raw.get("Readiness") else None))
        import_id=hashlib.sha256((checksum+stream).encode()).hexdigest();identity=report_identity or path.name
        with self.engine.db.transaction() as conn:conn.execute("INSERT INTO m2d_imports VALUES(?,?,?,?,?,?,?,?)",(import_id,identity,checksum,stream,None,len(rows),0,now()))
        imported=0
        for row in rows:
            self.engine.observe(stream=stream,source_id="jobqueue-import",provider_id=row.job_key,url=row.exact_vacancy_url,company=row.company,title=row.title,evidence={"report_identity":identity,"imported_status_claim":row.imported_status_claim,"revalidation_required":True});imported+=1
        with self.engine.db.transaction() as conn:conn.execute("UPDATE m2d_imports SET rows_imported=? WHERE import_id=?",(imported,import_id))
        return {"import_id":import_id,"rows_imported":imported,"duplicate":False}
